import json, re, requests, subprocess, pathlib
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

REPO_DIR=pathlib.Path("/tmp/inec-polling-units")
subprocess.run(["git","clone","--depth","1","https://github.com/mykeels/inec-polling-units.git",str(REPO_DIR)],check=True,stdout=subprocess.DEVNULL)
paths=list(REPO_DIR.glob("states/04-anambra/lgas/*/wards/*/units/index.json"))
assert len(paths)==326, f"Expected 326 ward unit files, got {len(paths)}"

rows=[]
for path in paths:
    units=json.loads(path.read_text(encoding="utf-8"))
    for u in units:
            d=u.get("delimitation","").replace("/","-")
            if not re.fullmatch(r"04-\d{2}-\d{2}-\d{3}",d):
                continue
            remark=u.get("remark","").strip().upper()
            status="New" if remark=="NEW PU" else "Existing" if remark=="EXISTING PU" else remark
            rows.append((d,u.get("local_government_id",""),u.get("local_government_name","").strip(),d.split("-")[2],u.get("ward_name","").strip(),d.split("-")[3],u.get("name","").strip(),status))

assert len({x[0] for x in rows})==len(rows), "Duplicate PU codes"
print("LGA_COUNTS", sorted({k:sum(1 for x in rows if x[1]==k) for k in {x[1] for x in rows}}.items(), key=lambda z:int(z[0])))
assert len(rows)==5720, f"Expected 5720 PUs, got {len(rows)}"
assert len({(x[1],x[3]) for x in rows})==326, "Expected 326 wards"
assert len({x[1] for x in rows})==21, "Expected 21 LGAs"
rows.sort(key=lambda x:tuple(map(int,x[0].split("-"))))

def pretty(s):
    return re.sub(r"\s+"," ",s.title()).strip()

doc=Document()
sec=doc.sections[0]
sec.orientation=WD_ORIENT.LANDSCAPE
sec.page_width,sec.page_height=sec.page_height,sec.page_width
sec.top_margin=Inches(.45); sec.bottom_margin=Inches(.45); sec.left_margin=Inches(.35); sec.right_margin=Inches(.35)
p=doc.add_paragraph(); p.alignment=1
x=p.add_run("ANAMBRA STATE — COMPLETE POLLING UNIT DIRECTORY"); x.bold=True; x.font.size=Pt(15)
p=doc.add_paragraph(); p.alignment=1
x=p.add_run("5,720 Polling Units • 326 Wards • 21 Local Government Areas"); x.bold=True; x.font.size=Pt(10)
p=doc.add_paragraph("Source: INEC-derived Anambra polling-unit records from the mykeels INEC Polling Units repository, which preserves the INEC delimitation code, polling-unit name, ward/LGA hierarchy and Existing/New PU remark. Specific Location / Address follows the recorded polling-unit/facility name, matching the established directory format; it is not presented as a separate street address where none is published.")
p.runs[0].font.size=Pt(8)

headers=["S/N","LGA Code","LGA","Ward Code","Ward","Polling Unit Code","Polling Unit","Specific Location / Address","Status"]
widths=[.38,.58,.9,.62,1.0,1.05,2.15,3.65,1.0]
t=doc.add_table(rows=1,cols=9); t.style="Table Grid"; t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
for i,h in enumerate(headers):
    c=t.rows[0].cells[i]; c.text=h; c.width=Inches(widths[i]); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for run in c.paragraphs[0].runs:
        run.bold=True; run.font.size=Pt(7)

for i,(full,lgacode,lga,wcode,ward,puc,pu,status) in enumerate(rows,1):
    vals=[str(i),f"04-{int(lgacode):02d}",pretty(lga),wcode,pretty(ward),full,pu,pretty(pu),status]
    cs=t.add_row().cells
    for j,v in enumerate(vals):
        cs[j].text=v; cs[j].width=Inches(widths[j]); cs[j].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for p in cs[j].paragraphs:
            for run in p.runs:
                run.font.size=Pt(6.5)

trPr=t.rows[0]._tr.get_or_add_trPr()
hdr=OxmlElement("w:tblHeader"); hdr.set(qn("w:val"),"true"); trPr.append(hdr)
footer=doc.sections[0].footer.paragraphs[0]; footer.alignment=1
footer.add_run("Anambra State Polling Unit Directory • INEC-derived • 25 September 2026").font.size=Pt(7)
doc.save("Anambra_5720_Polling_Units_INEC_Locator_Directory.docx")
print("Generated exactly 5,720 PUs / 326 wards / 21 LGAs")
